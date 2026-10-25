import datetime
import json
import os
import queue
import re
import subprocess
import sys
import threading
import time
import tkinter as tk
import tkinter.font as tkfont
from concurrent.futures import ThreadPoolExecutor, as_completed
from functools import lru_cache
from io import BytesIO
from urllib.parse import urljoin, urlparse

import requests
from requests.adapters import HTTPAdapter
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Cm as word_Cm
from PIL import Image, ImageDraw, ImageFilter, ImageFont, ImageTk, UnidentifiedImageError
from pptx import Presentation
from pptx.util import Cm as ppt_Cm
from tkinter import filedialog, ttk

USER_AGENT = (
    "Mozilla/5.0 (iPhone; CPU iPhone OS 16_0 like Mac OS X) "
    "AppleWebKit/605.1.15 (KHTML, like Gecko) Version/16.0 "
    "Mobile/15E148 Safari/604.1"
)
REQUEST_TIMEOUT = (5, 12)
MAX_DOWNLOAD_WORKERS = 10
MAX_DOWNLOAD_RETRIES = 3
RETRY_BACKOFF_SECONDS = 0.5
MAX_FILENAME_LENGTH = 80
DEFAULT_PREFIX = "公众号素材"
# 1×1、2×2 之类的透明占位图/统计像素直接跳过
TINY_IMAGE_PIXELS = 16
# Word / PowerPoint 能直接插入的格式；其余（webp 等）统一转成 PNG
OFFICE_FORMATS = {"JPEG": "jpg", "PNG": "png", "GIF": "gif", "BMP": "bmp"}
WORD_PAGE_WIDTH_CM = 21.0
WORD_PAGE_HEIGHT_CM = 29.7
WORD_MARGIN_CM = 2.5
WORD_MAX_WIDTH_CM = WORD_PAGE_WIDTH_CM - WORD_MARGIN_CM * 2
WORD_MAX_HEIGHT_CM = WORD_PAGE_HEIGHT_CM - WORD_MARGIN_CM * 2 - 0.7
WORD_IMAGE_DPI = 96
DOWNLOAD_PROGRESS_SHARE = 0.85
CONFIG_PATH = os.path.join(
    os.environ.get("APPDATA") or os.path.expanduser("~"), "微信公众号素材整理", "config.json"
)
CAPTCHA_HINT = "微信要求人机验证，请稍后再试，或先在浏览器中打开该文章完成验证"
BLOCKED_PAGE_HINTS = {
    "环境异常": CAPTCHA_HINT,
    "完成验证后即可继续访问": CAPTCHA_HINT,
    "该内容已被发布者删除": "文章已被发布者删除",
    "此内容因违规无法查看": "文章因违规无法查看",
    "此内容被投诉且经审核涉嫌侵权": "文章因侵权投诉无法查看",
}
WECHAT_IMAGE_PATTERN = re.compile(r"https?:(?:\\?/){2}mmbiz\.qpic\.cn\\?/[^\"'\s<>()]+")

# 界面配色：与 Oracle-Sync 桌宠菜单保持一致的白底卡片风格
WINDOW_BG = "#F5F6F8"
CARD_BG = "#FFFFFF"
BORDER = "#E4E6EB"
DIVIDER = "#ECEEF1"
HOVER_BG = "#F1F3F6"
PRESSED_BG = "#E8EBF0"
INFO_BG = "#F7F8FA"
TEXT = "#1F2329"
MUTED = "#6B7280"
PLACEHOLDER = "#A0A6B1"
DISABLED_TEXT = "#B8BDC6"
ICON = "#6B7280"
ACCENT = "#3B82F6"
ACCENT_HOVER = "#2563EB"
ACCENT_PRESSED = "#1D4ED8"
ACCENT_DISABLED = "#B7D0FB"
ACCENT_SOFT = "#F0F6FF"
SELECTION_BG = "#DBEAFE"
SUCCESS = "#15803D"
WARNING = "#B45309"
DANGER = "#D32F2F"
SHADOW_COLOR = "#111827"
SHADOW_ALPHA = 34

UI_FONT = "Microsoft YaHei UI"
ICON_FONT = "Segoe MDL2 Assets"
ICON_FONT_FILE = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", "segmdl2.ttf")

GLYPH_PICTURE = "\uE8B9"
GLYPH_LINK = "\uE71B"
GLYPH_RENAME = "\uE8AC"
GLYPH_FOLDER = "\uF12B"
GLYPH_FOLDER_OPEN = "\uED25"
GLYPH_PASTE = "\uE77F"
GLYPH_DOCUMENT = "\uE8A5"
GLYPH_SLIDESHOW = "\uE786"
GLYPH_PLAY = "\uE768"
GLYPH_CANCEL = "\uE711"
GLYPH_SETTINGS = "\uE713"
GLYPH_LIST = "\uE8FD"
GLYPH_CHECKBOX = "\uE739"
GLYPH_CHECKBOX_ON = "\uE73A"
LOG_GLYPHS = {"info": "\uE946", "success": "\uE930", "warning": "\uE7BA", "error": "\uE783"}
LOG_COLORS = {"info": TEXT, "success": SUCCESS, "warning": WARNING, "error": DANGER}
LOG_ICON_COLORS = {"info": PLACEHOLDER, "success": SUCCESS, "warning": WARNING, "error": DANGER}

UI_SCALE = 1.0


def px(value: float) -> int:
    return int(round(value * UI_SCALE))


class Cancelled(Exception):
    pass


class SkippedImage(Exception):
    pass


def sanitize_filename(name: str) -> str:
    cleaned = re.sub(r'[\\/*?:"<>|\r\n\t]', "_", name or "")
    cleaned = cleaned.replace("..", "_").strip(" ._")
    if not cleaned:
        cleaned = DEFAULT_PREFIX
    return cleaned[:MAX_FILENAME_LENGTH]


def is_http_url(value: str) -> bool:
    parsed = urlparse(value)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def post_event(status_queue: queue.Queue, event_type: str, **payload: object) -> None:
    event = {"type": event_type}
    event.update(payload)
    status_queue.put(event)


def log_status(status_queue: queue.Queue, message: str, level: str = "info") -> None:
    post_event(status_queue, "log", message=message, level=level)


def set_stage(status_queue: queue.Queue, label: str) -> None:
    post_event(status_queue, "stage", label=label)


def set_progress(status_queue: queue.Queue, fraction: float) -> None:
    post_event(status_queue, "progress", value=max(0.0, min(1.0, fraction)))


def open_in_file_manager(path: str) -> None:
    try:
        if sys.platform.startswith("win"):
            os.startfile(path)
        elif sys.platform == "darwin":
            subprocess.Popen(["open", path])
        else:
            subprocess.Popen(["xdg-open", path])
    except OSError:
        pass


def load_config() -> dict:
    try:
        with open(CONFIG_PATH, encoding="utf-8") as file_obj:
            data = json.load(file_obj)
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def save_config(data: dict) -> None:
    try:
        os.makedirs(os.path.dirname(CONFIG_PATH), exist_ok=True)
        with open(CONFIG_PATH, "w", encoding="utf-8") as file_obj:
            json.dump(data, file_obj, ensure_ascii=False, indent=2)
    except OSError:
        pass


def default_save_dir() -> str:
    # 桌面可能被重定向到其他盘，优先读系统登记的真实路径
    if sys.platform.startswith("win"):
        try:
            import winreg

            with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                r"Software\Microsoft\Windows\CurrentVersion\Explorer\Shell Folders",
            ) as key:
                desktop = winreg.QueryValueEx(key, "Desktop")[0]
            if os.path.isdir(desktop):
                return desktop
        except OSError:
            pass
    desktop = os.path.join(os.path.expanduser("~"), "Desktop")
    return desktop if os.path.isdir(desktop) else ""


def create_session_folder(base_path: str, name: str) -> str:
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    folder_path = os.path.join(base_path, f"{name}_{timestamp}")
    os.makedirs(folder_path, exist_ok=True)
    return folder_path


def build_session() -> requests.Session:
    session = requests.Session()
    # 微信公众号为国内服务器，避免因系统代理/梯子导致流量绕行海外引起 20 秒超时与下载降速
    session.trust_env = False
    adapter = HTTPAdapter(pool_connections=16, pool_maxsize=16)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    session.headers.update({"User-Agent": USER_AGENT})
    return session


def extract_title(soup: BeautifulSoup) -> str:
    meta = soup.find("meta", attrs={"property": "og:title"})
    if meta and meta.get("content", "").strip():
        return meta["content"].strip()
    node = soup.select_one("#activity-name") or soup.select_one(".rich_media_title")
    if node and node.get_text(strip=True):
        return node.get_text(strip=True)
    if soup.title and soup.title.string:
        return soup.title.string.strip()
    return ""


def detect_blocked_page(soup: BeautifulSoup) -> str:
    text = soup.get_text(" ", strip=True)
    for keyword, hint in BLOCKED_PAGE_HINTS.items():
        if keyword in text:
            return hint
    return ""


def extract_image_urls(page_url: str, soup: BeautifulSoup) -> list[str]:
    # 只取正文区域，避免把头像、二维码、推荐阅读的封面一起下载
    container = soup.select_one("#js_content") or soup.select_one(".rich_media_content") or soup
    attrs = ("data-src", "data-actualsrc", "data-original", "data-lazy-src", "src")
    found_urls: list[str] = []
    seen: set[str] = set()

    for img in container.find_all("img"):
        for attr in attrs:
            raw_url = img.get(attr)
            if not raw_url:
                continue
            image_url = urljoin(page_url, raw_url.strip())
            if not is_http_url(image_url):
                continue
            if image_url not in seen:
                seen.add(image_url)
                found_urls.append(image_url)
            break

    return found_urls


def extract_script_image_urls(raw_html: str) -> list[str]:
    # “图片消息”类文章没有 <img>，图片地址只存在于页面脚本里
    found_urls: list[str] = []
    seen: set[str] = set()
    for match in WECHAT_IMAGE_PATTERN.finditer(raw_html):
        image_url = (
            match.group(0)
            .replace("\\/", "/")
            .replace("\\x26amp;", "&")
            .replace("\\x26", "&")
            .replace("&amp;", "&")
            .rstrip("\\")
        )
        # 同一张图常以不同尺寸参数出现多次，按去掉尺寸段后的路径去重
        key = urlparse(image_url).path.rsplit("/", 1)[0]
        if key not in seen:
            seen.add(key)
            found_urls.append(image_url)
    return found_urls


def fetch_article(session: requests.Session, article_url: str) -> tuple[str, list[str], str]:
    response = session.get(article_url, timeout=REQUEST_TIMEOUT)
    response.raise_for_status()
    soup = BeautifulSoup(response.content, "lxml")
    image_urls = extract_image_urls(article_url, soup)
    if not image_urls:
        image_urls = extract_script_image_urls(response.content.decode("utf-8", "replace"))
    blocked_reason = "" if image_urls else detect_blocked_page(soup)
    return extract_title(soup), image_urls, blocked_reason


def save_image_bytes(image_bytes: bytes, save_folder: str, index: int) -> str:
    try:
        with Image.open(BytesIO(image_bytes)) as image:
            width, height = image.size
            if width * height <= TINY_IMAGE_PIXELS:
                raise SkippedImage(f"{width}×{height} 占位图")
            base_path = os.path.join(save_folder, f"{index:03d}")
            image_format = (image.format or "").upper()
            if image_format in OFFICE_FORMATS:
                file_path = f"{base_path}.{OFFICE_FORMATS[image_format]}"
                with open(file_path, "wb") as file_obj:
                    file_obj.write(image_bytes)
            else:
                file_path = f"{base_path}.png"
                has_alpha = image.mode in ("RGBA", "LA") or (
                    image.mode == "P" and "transparency" in image.info
                )
                image.convert("RGBA" if has_alpha else "RGB").save(file_path, "PNG")
            return file_path
    except (UnidentifiedImageError, Image.DecompressionBombError) as exc:
        raise SkippedImage("不是可识别的图片") from exc


def download_single_image(
    session: requests.Session,
    image_url: str,
    save_folder: str,
    index: int,
    cancel_event: threading.Event,
) -> str:
    last_error: Exception | None = None
    for attempt in range(1, MAX_DOWNLOAD_RETRIES + 1):
        if cancel_event.is_set():
            raise Cancelled()
        try:
            response = session.get(image_url, timeout=REQUEST_TIMEOUT)
            response.raise_for_status()
            return save_image_bytes(response.content, save_folder, index)
        except requests.RequestException as exc:
            last_error = exc
            status = exc.response.status_code if exc.response is not None else 0
            if 400 <= status < 500 and status != 429:
                # 404/403 之类重试也没用
                raise
            if attempt < MAX_DOWNLOAD_RETRIES:
                if cancel_event.wait(RETRY_BACKOFF_SECONDS * attempt):
                    raise Cancelled()
    assert last_error is not None
    raise last_error


def download_images(
    session: requests.Session,
    image_urls: list[str],
    save_folder: str,
    status_queue: queue.Queue,
    cancel_event: threading.Event,
) -> list[str]:
    downloaded: list[tuple[int, str]] = []
    completed_count = 0
    total_count = len(image_urls)
    set_stage(status_queue, f"下载图片 0/{total_count}")
    set_progress(status_queue, 0)

    executor = ThreadPoolExecutor(max_workers=MAX_DOWNLOAD_WORKERS)
    try:
        future_map = {
            executor.submit(
                download_single_image, session, image_url, save_folder, index, cancel_event
            ): (index, image_url)
            for index, image_url in enumerate(image_urls, start=1)
        }
        for future in as_completed(future_map):
            if cancel_event.is_set():
                break
            index, image_url = future_map[future]
            try:
                file_path = future.result()
                downloaded.append((index, file_path))
            except Cancelled:
                break
            except SkippedImage as exc:
                log_status(status_queue, f"第 {index} 张已跳过：{exc}")
            # requests 的异常继承自 OSError，必须先于 OSError 捕获
            except requests.RequestException as exc:
                log_status(status_queue, f"第 {index} 张下载失败：{exc}", "warning")
            except OSError as exc:
                log_status(status_queue, f"第 {index} 张保存失败：{exc}", "warning")
            completed_count += 1
            set_stage(status_queue, f"下载图片 {completed_count}/{total_count}")
            set_progress(status_queue, completed_count / total_count * DOWNLOAD_PROGRESS_SHARE)
    finally:
        # 取消时不等待仍在进行的请求，让界面立即恢复
        executor.shutdown(wait=not cancel_event.is_set(), cancel_futures=True)

    downloaded.sort()
    image_paths = [path for _, path in downloaded]
    if cancel_event.is_set():
        log_status(status_queue, f"已取消，已保留 {len(image_paths)} 张图片", "warning")
    else:
        log_status(status_queue, f"图片下载完成：成功 {len(image_paths)} / 共 {total_count} 张")
    return image_paths


def word_width_for(image_path: str) -> float:
    try:
        with Image.open(image_path) as image:
            width_px, height_px = image.size
    except (UnidentifiedImageError, OSError):
        return WORD_MAX_WIDTH_CM
    natural_width = width_px / WORD_IMAGE_DPI * 2.54
    natural_height = height_px / WORD_IMAGE_DPI * 2.54
    # 不放大小图；长图同时按页面可用高度缩小，避免超出页面被裁掉
    scale = min(1.0, WORD_MAX_WIDTH_CM / natural_width, WORD_MAX_HEIGHT_CM / natural_height)
    return natural_width * scale


def generate_word_document(name: str, image_paths: list[str], folder: str, status_queue: queue.Queue) -> None:
    if not image_paths:
        return

    set_stage(status_queue, "生成 Word")
    log_status(status_queue, "正在生成 Word 文档")
    document = Document()
    section = document.sections[0]
    section.page_width = word_Cm(WORD_PAGE_WIDTH_CM)
    section.page_height = word_Cm(WORD_PAGE_HEIGHT_CM)
    section.left_margin = section.right_margin = word_Cm(WORD_MARGIN_CM)
    section.top_margin = section.bottom_margin = word_Cm(WORD_MARGIN_CM)
    for image_path in image_paths:
        try:
            document.add_picture(image_path, width=word_Cm(word_width_for(image_path)))
        except Exception as exc:
            log_status(status_queue, f"Word 跳过 {os.path.basename(image_path)}：{exc}", "warning")

    file_path = os.path.join(folder, f"{name}.docx")
    document.save(file_path)
    log_status(status_queue, f"已保存 {os.path.basename(file_path)}", "success")


def fit_image_to_slide(image_width: int, image_height: int) -> tuple[float, float, float, float]:
    slide_width = 25.4
    slide_height = 19.05
    margin = 1.0
    max_width = slide_width - margin * 2
    max_height = slide_height - margin * 2
    scale = min(max_width / image_width, max_height / image_height)
    final_width = image_width * scale
    final_height = image_height * scale
    left = (slide_width - final_width) / 2
    top = (slide_height - final_height) / 2
    return left, top, final_width, final_height


def generate_ppt_presentation(name: str, image_paths: list[str], folder: str, status_queue: queue.Queue) -> None:
    if not image_paths:
        return

    set_stage(status_queue, "生成 PowerPoint")
    log_status(status_queue, "正在生成 PowerPoint")
    presentation = Presentation()
    while presentation.slides:
        rel_id = presentation.slides._sldIdLst[0].rId
        presentation.part.drop_rel(rel_id)
        del presentation.slides._sldIdLst[0]

    for image_path in image_paths:
        try:
            with Image.open(image_path) as image:
                width, height = image.size
            left, top, final_width, final_height = fit_image_to_slide(width, height)
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(
                image_path,
                ppt_Cm(left),
                ppt_Cm(top),
                width=ppt_Cm(final_width),
                height=ppt_Cm(final_height),
            )
        except Exception as exc:
            log_status(status_queue, f"PPT 跳过 {os.path.basename(image_path)}：{exc}", "warning")

    file_path = os.path.join(folder, f"{name}.pptx")
    presentation.save(file_path)
    log_status(status_queue, f"已保存 {os.path.basename(file_path)}", "success")


def run_export(
    article_url: str,
    prefix: str,
    save_dir: str,
    want_word: bool,
    want_ppt: bool,
    status_queue: queue.Queue,
    cancel_event: threading.Event,
) -> None:
    """后台线程入口：只通过 status_queue 与界面通信，结束时一定发出 done 事件。"""
    started = time.monotonic()
    output_folder = ""
    outcome = "failed"
    try:
        session = build_session()
        set_stage(status_queue, "获取文章页面")
        log_status(status_queue, "正在获取文章页面")
        try:
            title, image_urls, blocked_reason = fetch_article(session, article_url)
        except requests.RequestException as exc:
            log_status(status_queue, f"文章页面获取失败：{exc}", "error")
            return

        if title:
            log_status(status_queue, f"文章标题：{title}")
        if not image_urls:
            log_status(status_queue, blocked_reason or "正文中没有找到图片", "error")
            return
        if cancel_event.is_set():
            outcome = "cancelled"
            return

        log_status(status_queue, f"共发现 {len(image_urls)} 张图片")
        name = sanitize_filename(prefix or title or DEFAULT_PREFIX)
        output_folder = create_session_folder(save_dir, name)
        log_status(status_queue, f"输出目录：{output_folder}")

        image_paths = download_images(session, image_urls, output_folder, status_queue, cancel_event)
        if cancel_event.is_set():
            outcome = "cancelled"
            return
        if not image_paths:
            log_status(status_queue, "没有成功下载的图片，未生成文档", "error")
            return

        if want_word:
            set_progress(status_queue, DOWNLOAD_PROGRESS_SHARE + 0.03)
            generate_word_document(name, image_paths, output_folder, status_queue)
        if want_ppt and not cancel_event.is_set():
            set_progress(status_queue, DOWNLOAD_PROGRESS_SHARE + 0.09)
            generate_ppt_presentation(name, image_paths, output_folder, status_queue)

        outcome = "cancelled" if cancel_event.is_set() else "success"
        if outcome == "success":
            set_progress(status_queue, 1.0)
            log_status(status_queue, "全部完成", "success")
    except Exception as exc:
        log_status(status_queue, f"发生未预期错误：{exc}", "error")
    finally:
        post_event(
            status_queue,
            "done",
            outcome=outcome,
            folder=output_folder,
            elapsed=time.monotonic() - started,
        )


# ---------------------------------------------------------------- 界面组件


@lru_cache(maxsize=256)
def render_round_rect(
    width: int,
    height: int,
    radius: int,
    fill: str,
    outline: str | None,
    background: str,
    shadow: int = 0,
) -> Image.Image:
    """抗锯齿圆角矩形；shadow > 0 时四周留出该宽度的边距绘制柔和阴影。"""
    width, height = max(width, 1), max(height, 1)
    canvas = Image.new("RGB", (width, height), background)
    box_width, box_height = width - shadow * 2, height - shadow * 2
    if box_width < 4 or box_height < 4:
        return canvas

    if shadow:
        mask = Image.new("L", (width, height), 0)
        offset = max(1, shadow // 5)
        ImageDraw.Draw(mask).rounded_rectangle(
            (shadow, shadow + offset, width - shadow - 1, height - shadow - 1 + offset),
            radius=radius,
            fill=SHADOW_ALPHA,
        )
        mask = mask.filter(ImageFilter.GaussianBlur(shadow / 2.6))
        canvas = Image.composite(Image.new("RGB", (width, height), SHADOW_COLOR), canvas, mask)

    supersample = 3
    shape = Image.new("RGBA", (box_width * supersample, box_height * supersample), (0, 0, 0, 0))
    ImageDraw.Draw(shape).rounded_rectangle(
        (0, 0, box_width * supersample - 1, box_height * supersample - 1),
        radius=radius * supersample,
        fill=fill,
        outline=outline,
        width=supersample if outline else 0,
    )
    shape = shape.resize((box_width, box_height), Image.LANCZOS)
    canvas.paste(shape, (shadow, shadow), shape)
    return canvas


def recolor_tree(widget: tk.Misc, old_bg: str, new_bg: str) -> None:
    for child in widget.winfo_children():
        try:
            if str(child.cget("bg")).upper() == old_bg.upper():
                child.configure(bg=new_bg)
        except tk.TclError:
            pass
        recolor_tree(child, old_bg, new_bg)


def bind_tree(widget: tk.Misc, sequence: str, callback) -> None:
    widget.bind(sequence, callback, add="+")
    for child in widget.winfo_children():
        bind_tree(child, sequence, callback)


def elide_middle(text: str, font: tkfont.Font, max_width: int) -> str:
    if max_width <= 0 or font.measure(text) <= max_width:
        return text
    low, high = 0, len(text)
    while low < high:
        keep = (low + high + 1) // 2
        candidate = text[: keep // 2] + "…" + text[len(text) - (keep - keep // 2):]
        if font.measure(candidate) <= max_width:
            low = keep
        else:
            high = keep - 1
    return text[: low // 2] + "…" + text[len(text) - (low - low // 2):]


class RoundedBox(tk.Canvas):
    """圆角卡片容器：背景用 PIL 抗锯齿绘制，子控件放在 self.inner 里。"""

    def __init__(
        self,
        parent: tk.Misc,
        *,
        parent_bg: str,
        fill: str,
        outline: str | None = None,
        radius: int = 10,
        padding: tuple[int, int] = (0, 0),
        shadow: int = 0,
        height: int | None = None,
    ) -> None:
        super().__init__(
            parent,
            bg=parent_bg,
            highlightthickness=0,
            bd=0,
            width=1,
            height=px(height) if height else 1,
        )
        self.parent_bg = parent_bg
        self.fill = fill
        self.outline = outline
        self.radius = px(radius)
        self.shadow = px(shadow) if shadow else 0
        self.offset_x = self.shadow + px(padding[0])
        self.offset_y = self.shadow + px(padding[1])
        self.fit_content = height is None
        self._size = (0, 0)
        self._photo: ImageTk.PhotoImage | None = None

        self.inner = tk.Frame(self, bg=fill)
        self._image_item = self.create_image(0, 0, anchor="nw")
        self._window_item = self.create_window(self.offset_x, self.offset_y, window=self.inner, anchor="nw")
        self.bind("<Configure>", self._on_configure)
        if self.fit_content:
            # 画布太小时 inner 不会被映射、也收不到 <Configure>，首次高度要主动同步
            self.inner.bind("<Configure>", self._sync_height)
            self.after_idle(self._sync_height)

    def _on_configure(self, event: tk.Event) -> None:
        self.itemconfigure(self._window_item, width=max(1, event.width - self.offset_x * 2))
        if not self.fit_content:
            self.itemconfigure(self._window_item, height=max(1, event.height - self.offset_y * 2))
        self._size = (event.width, event.height)
        self._redraw()

    def _sync_height(self, _event: tk.Event | None = None) -> None:
        needed = self.inner.winfo_reqheight() + self.offset_y * 2
        if int(float(self.cget("height"))) != needed:
            self.configure(height=needed)

    def _redraw(self) -> None:
        width, height = self._size
        if width < 2 or height < 2:
            return
        image = render_round_rect(
            width, height, self.radius, self.fill, self.outline, self.parent_bg, self.shadow
        )
        self._photo = ImageTk.PhotoImage(image)
        self.itemconfigure(self._image_item, image=self._photo)

    def set_style(self, *, fill: str | None = None, outline: str | None = "") -> None:
        new_fill = fill or self.fill
        new_outline = self.outline if outline == "" else outline
        if new_fill == self.fill and new_outline == self.outline:
            return
        if new_fill != self.fill:
            self.inner.configure(bg=new_fill)
            recolor_tree(self.inner, self.fill, new_fill)
        self.fill, self.outline = new_fill, new_outline
        self._redraw()


class IconButton(tk.Canvas):
    # 每种状态：(填充, 描边, 文字色, 图标色)；填充为 None 时与父容器同色
    VARIANTS = {
        "primary": {
            "normal": (ACCENT, None, "#FFFFFF", "#FFFFFF"),
            "hover": (ACCENT_HOVER, None, "#FFFFFF", "#FFFFFF"),
            "pressed": (ACCENT_PRESSED, None, "#FFFFFF", "#FFFFFF"),
            "disabled": (ACCENT_DISABLED, None, "#FFFFFF", "#FFFFFF"),
        },
        "secondary": {
            "normal": (CARD_BG, BORDER, TEXT, ICON),
            "hover": (HOVER_BG, BORDER, TEXT, ACCENT),
            "pressed": (PRESSED_BG, BORDER, TEXT, ACCENT),
            "disabled": (CARD_BG, DIVIDER, DISABLED_TEXT, DISABLED_TEXT),
        },
        "ghost": {
            "normal": (None, None, MUTED, ICON),
            "hover": (HOVER_BG, None, TEXT, ACCENT),
            "pressed": (PRESSED_BG, None, TEXT, ACCENT),
            "disabled": (None, None, DISABLED_TEXT, DISABLED_TEXT),
        },
    }

    def __init__(
        self,
        parent: tk.Misc,
        *,
        text: str,
        icon: str,
        command,
        parent_bg: str,
        font: tkfont.Font,
        icon_font: tkfont.Font,
        variant: str = "secondary",
        height: int = 38,
        padding_x: int = 16,
    ) -> None:
        self.text = text
        self.icon = icon
        self.command = command
        self.parent_bg = parent_bg
        self.text_font = font
        self.icon_font = icon_font
        self.styles = self.VARIANTS[variant]
        self.padding_x = px(padding_x)
        self.gap = px(8)
        self.enabled = True
        self.hover = False
        self.pressed = False
        self._photo: ImageTk.PhotoImage | None = None
        super().__init__(
            parent,
            bg=parent_bg,
            highlightthickness=0,
            bd=0,
            width=self._natural_width(),
            height=px(height),
            cursor="hand2",
        )
        self._image_item = self.create_image(0, 0, anchor="nw")
        self._icon_item = self.create_text(0, 0, text=icon, font=icon_font, anchor="w")
        self._text_item = self.create_text(0, 0, text=text, font=font, anchor="w")
        self.bind("<Configure>", lambda _event: self._redraw())
        self.bind("<Enter>", self._on_enter)
        self.bind("<Leave>", self._on_leave)
        self.bind("<ButtonPress-1>", self._on_press)
        self.bind("<ButtonRelease-1>", self._on_release)

    def _content_width(self) -> int:
        width = self.icon_font.measure(self.icon) if self.icon else 0
        if self.icon and self.text:
            width += self.gap
        if self.text:
            width += self.text_font.measure(self.text)
        return width

    def _natural_width(self) -> int:
        return self._content_width() + self.padding_x * 2

    def _state(self) -> str:
        if not self.enabled:
            return "disabled"
        if self.pressed and self.hover:
            return "pressed"
        return "hover" if self.hover else "normal"

    def _redraw(self) -> None:
        width, height = self.winfo_width(), self.winfo_height()
        if width < 4 or height < 4:
            return
        fill, outline, text_color, icon_color = self.styles[self._state()]
        image = render_round_rect(width, height, px(8), fill or self.parent_bg, outline, self.parent_bg)
        self._photo = ImageTk.PhotoImage(image)
        self.itemconfigure(self._image_item, image=self._photo)

        x = (width - self._content_width()) / 2
        y = height / 2
        if self.icon:
            self.coords(self._icon_item, x, y + 1)
            self.itemconfigure(self._icon_item, fill=icon_color)
            x += self.icon_font.measure(self.icon) + self.gap
        self.coords(self._text_item, x, y)
        self.itemconfigure(self._text_item, fill=text_color)

    def _on_enter(self, _event: tk.Event) -> None:
        self.hover = True
        self._redraw()

    def _on_leave(self, _event: tk.Event) -> None:
        self.hover = False
        self._redraw()

    def _on_press(self, _event: tk.Event) -> None:
        if self.enabled:
            self.pressed = True
            self._redraw()

    def _on_release(self, _event: tk.Event) -> None:
        fire = self.enabled and self.pressed and self.hover
        self.pressed = False
        self._redraw()
        if fire and self.command:
            self.command()

    def set_enabled(self, enabled: bool) -> None:
        self.enabled = enabled
        self.configure(cursor="hand2" if enabled else "arrow")
        self._redraw()

    def set_text(self, text: str) -> None:
        self.text = text
        self.itemconfigure(self._text_item, text=text)
        self.configure(width=self._natural_width())
        self._redraw()


class InputField(RoundedBox):
    def __init__(
        self,
        parent: tk.Misc,
        *,
        parent_bg: str,
        icon: str,
        placeholder: str,
        font: tkfont.Font,
        icon_font: tkfont.Font,
    ) -> None:
        super().__init__(
            parent,
            parent_bg=parent_bg,
            fill=CARD_BG,
            outline=BORDER,
            radius=8,
            padding=(12, 1),
            height=40,
        )
        self.focused = False
        self.has_error = False
        self.var = tk.StringVar()

        tk.Label(self.inner, text=icon, font=icon_font, bg=CARD_BG, fg=ICON).pack(
            side="left", padx=(0, px(10))
        )
        self.actions = tk.Frame(self.inner, bg=CARD_BG)
        self.actions.pack(side="right")
        self.entry = tk.Entry(
            self.inner,
            textvariable=self.var,
            font=font,
            bg=CARD_BG,
            fg=TEXT,
            relief="flat",
            bd=0,
            highlightthickness=0,
            insertbackground=TEXT,
            selectbackground=SELECTION_BG,
            selectforeground=TEXT,
        )
        self.entry.pack(side="left", fill="x", expand=True)
        self.placeholder = tk.Label(
            self.inner, text=placeholder, font=font, bg=CARD_BG, fg=PLACEHOLDER, bd=0, padx=0, cursor="xterm"
        )
        self.placeholder.bind("<Button-1>", lambda _event: self.entry.focus_set())

        self.entry.bind("<FocusIn>", lambda _event: self._set_focus(True), add="+")
        self.entry.bind("<FocusOut>", lambda _event: self._set_focus(False), add="+")
        self.var.trace_add("write", lambda *_args: self._on_change())
        self._build_context_menu(font)
        self._on_change()

    def _build_context_menu(self, font: tkfont.Font) -> None:
        menu = tk.Menu(self, tearoff=False, font=font, bg=CARD_BG, fg=TEXT,
                       activebackground=HOVER_BG, activeforeground=TEXT, bd=0)
        menu.add_command(label="剪切", command=lambda: self.entry.event_generate("<<Cut>>"))
        menu.add_command(label="复制", command=lambda: self.entry.event_generate("<<Copy>>"))
        menu.add_command(label="粘贴", command=lambda: self.entry.event_generate("<<Paste>>"))
        menu.add_separator()
        menu.add_command(label="全选", command=lambda: self.entry.select_range(0, "end"))
        menu.add_command(label="清空", command=lambda: self.var.set(""))

        def popup(event: tk.Event) -> None:
            self.entry.focus_set()
            menu.tk_popup(event.x_root, event.y_root)

        self.entry.bind("<Button-3>", popup)

    def add_action(self, button_factory) -> None:
        button_factory(self.actions).pack(side="left")

    def _on_change(self) -> None:
        if self.var.get():
            self.placeholder.place_forget()
        else:
            self.placeholder.place(in_=self.entry, x=px(2), rely=0.5, anchor="w")
        if self.has_error:
            self.set_error(False)

    def _set_focus(self, focused: bool) -> None:
        self.focused = focused
        self._refresh_outline()

    def _refresh_outline(self) -> None:
        if self.has_error:
            self.set_style(outline=DANGER)
        else:
            self.set_style(outline=ACCENT if self.focused else BORDER)

    def set_error(self, has_error: bool) -> None:
        self.has_error = has_error
        self._refresh_outline()

    def get(self) -> str:
        return self.var.get().strip()


class OptionTile(RoundedBox):
    def __init__(
        self,
        parent: tk.Misc,
        *,
        parent_bg: str,
        title: str,
        subtitle: str,
        icon: str,
        variable: tk.BooleanVar,
        fonts: dict[str, tkfont.Font],
    ) -> None:
        super().__init__(
            parent,
            parent_bg=parent_bg,
            fill=CARD_BG,
            outline=BORDER,
            radius=8,
            padding=(14, 1),
            height=58,
        )
        self.variable = variable
        self.hover = False
        self.has_error = False

        self.icon_label = tk.Label(self.inner, text=icon, font=fonts["icon_large"], bg=CARD_BG, fg=ICON)
        self.icon_label.pack(side="left", padx=(0, px(12)))
        texts = tk.Frame(self.inner, bg=CARD_BG)
        texts.pack(side="left", fill="x", expand=True)
        tk.Label(texts, text=title, font=fonts["body"], bg=CARD_BG, fg=TEXT, anchor="w").pack(fill="x")
        tk.Label(texts, text=subtitle, font=fonts["small"], bg=CARD_BG, fg=MUTED, anchor="w").pack(fill="x")
        self.check_label = tk.Label(self.inner, font=fonts["icon"], bg=CARD_BG)
        self.check_label.pack(side="right")

        for sequence in ("<Enter>", "<Leave>"):
            bind_tree(self, sequence, lambda _event: self.after(1, self._check_hover))
        bind_tree(self, "<ButtonRelease-1>", self._toggle)
        self.configure(cursor="hand2")
        for widget in (self.inner, texts, self.icon_label, self.check_label, *texts.winfo_children()):
            widget.configure(cursor="hand2")
        variable.trace_add("write", lambda *_args: self._refresh())
        self._refresh()

    def _check_hover(self) -> None:
        try:
            pointer = self.winfo_containing(*self.winfo_pointerxy())
        except (tk.TclError, KeyError):
            pointer = None
        hover = pointer is not None and str(pointer).startswith(str(self))
        if hover != self.hover:
            self.hover = hover
            self._refresh()

    def _toggle(self, _event: tk.Event) -> None:
        if self.hover:
            self.variable.set(not self.variable.get())

    def set_error(self, has_error: bool) -> None:
        self.has_error = has_error
        self._refresh()

    def _refresh(self) -> None:
        checked = bool(self.variable.get())
        if checked and self.has_error:
            self.has_error = False
        if checked:
            fill, outline = ACCENT_SOFT, ACCENT
        else:
            fill = HOVER_BG if self.hover else CARD_BG
            outline = DANGER if self.has_error else BORDER
        self.set_style(fill=fill, outline=outline)
        self.icon_label.configure(fg=ACCENT if checked else ICON)
        self.check_label.configure(
            text=GLYPH_CHECKBOX_ON if checked else GLYPH_CHECKBOX,
            fg=ACCENT if checked else PLACEHOLDER,
        )


class ProgressBar(tk.Canvas):
    def __init__(self, parent: tk.Misc, *, parent_bg: str) -> None:
        super().__init__(parent, bg=parent_bg, highlightthickness=0, bd=0, width=1, height=px(6))
        self.parent_bg = parent_bg
        self.value = 0.0
        self.color = ACCENT
        self._photo: ImageTk.PhotoImage | None = None
        self._image_item = self.create_image(0, 0, anchor="nw")
        self.bind("<Configure>", lambda _event: self._redraw())

    def set(self, value: float, color: str | None = None) -> None:
        self.value = max(0.0, min(1.0, value))
        self.color = color or ACCENT
        self._redraw()

    def _redraw(self) -> None:
        width, height = self.winfo_width(), self.winfo_height()
        if width < 4 or height < 2:
            return
        image = render_round_rect(width, height, height // 2, DIVIDER, None, self.parent_bg).copy()
        filled = int(round(width * self.value))
        if filled > 0:
            bar = render_round_rect(max(filled, height), height, height // 2, self.color, None, DIVIDER)
            mask = render_round_rect(max(filled, height), height, height // 2, "#FFFFFF", None, "#000000")
            image.paste(bar, (0, 0), mask.convert("L"))
        self._photo = ImageTk.PhotoImage(image)
        self.itemconfigure(self._image_item, image=self._photo)


# ---------------------------------------------------------------- 主窗口


class MaterialApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("公众号素材整理")
        self.root.configure(bg=WINDOW_BG)
        width = px(720)
        height = min(px(900), root.winfo_screenheight() - px(80))
        left = max(0, (root.winfo_screenwidth() - width) // 2)
        top = max(0, (root.winfo_screenheight() - height) // 2 - px(20))
        self.root.geometry(f"{width}x{height}+{left}+{top}")
        self.root.minsize(px(600), min(px(700), height))

        config = load_config()
        saved_dir = str(config.get("save_dir") or "")
        self.save_dir = saved_dir if os.path.isdir(saved_dir) else default_save_dir()
        self.gen_word_var = tk.BooleanVar(value=bool(config.get("want_word", True)))
        self.gen_ppt_var = tk.BooleanVar(value=bool(config.get("want_ppt", True)))
        self.status_queue: queue.Queue = queue.Queue()
        self.cancel_event = threading.Event()
        self.is_running = False
        self.last_output_folder = ""
        self.log_is_placeholder = False

        self._init_fonts()
        self._setup_styles()
        self._set_window_icon()
        self._build_ui()
        self._bind_shortcuts()
        self._poll_status_queue()
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        self.root.after(50, self.url_field.entry.focus_set)

    def _init_fonts(self) -> None:
        self.fonts = {
            "title": tkfont.Font(family=UI_FONT, size=15, weight="bold"),
            "card_title": tkfont.Font(family=UI_FONT, size=11, weight="bold"),
            "body": tkfont.Font(family=UI_FONT, size=10),
            "label": tkfont.Font(family=UI_FONT, size=10),
            "small": tkfont.Font(family=UI_FONT, size=9),
            "button": tkfont.Font(family=UI_FONT, size=10),
            "button_bold": tkfont.Font(family=UI_FONT, size=10, weight="bold"),
            "log": tkfont.Font(family=UI_FONT, size=9),
            "icon": tkfont.Font(family=ICON_FONT, size=11),
            "icon_small": tkfont.Font(family=ICON_FONT, size=9),
            "icon_large": tkfont.Font(family=ICON_FONT, size=16),
            "icon_badge": tkfont.Font(family=ICON_FONT, size=17),
        }

    def _setup_styles(self) -> None:
        style = ttk.Style(self.root)
        style.theme_use("clam")
        style.layout(
            "Thin.Vertical.TScrollbar",
            [("Vertical.Scrollbar.trough", {"sticky": "ns", "children": [
                ("Vertical.Scrollbar.thumb", {"expand": "1", "sticky": "nswe"})
            ]})],
        )
        style.configure(
            "Thin.Vertical.TScrollbar",
            troughcolor=INFO_BG,
            background="#D7DBE1",
            bordercolor=INFO_BG,
            lightcolor="#D7DBE1",
            darkcolor="#D7DBE1",
            arrowsize=px(7),
            gripcount=0,
        )
        style.map("Thin.Vertical.TScrollbar", background=[("active", "#B9BFC8")])

    def _set_window_icon(self) -> None:
        try:
            size = 64
            image = render_round_rect(size, size, 14, ACCENT, None, "#FFFFFF").convert("RGBA")
            rounded = render_round_rect(size, size, 14, "#FFFFFF", None, "#000000").convert("L")
            image.putalpha(rounded)
            glyph_font = ImageFont.truetype(ICON_FONT_FILE, 34)
            ImageDraw.Draw(image).text((size / 2, size / 2), GLYPH_PICTURE, font=glyph_font, fill="#FFFFFF", anchor="mm")
            self._icon_photo = ImageTk.PhotoImage(image)
            self.root.iconphoto(True, self._icon_photo)
        except (OSError, tk.TclError):
            pass

    def _section_label(self, parent: tk.Misc, text: str, hint: str = "") -> None:
        row = tk.Frame(parent, bg=CARD_BG)
        row.pack(fill="x", pady=(px(14), px(6)))
        tk.Label(row, text=text, font=self.fonts["label"], bg=CARD_BG, fg=TEXT).pack(side="left")
        if hint:
            tk.Label(row, text=hint, font=self.fonts["small"], bg=CARD_BG, fg=MUTED).pack(side="right")

    def _card_header(self, parent: tk.Misc, icon: str, title: str) -> tk.Frame:
        row = tk.Frame(parent, bg=CARD_BG)
        row.pack(fill="x")
        tk.Label(row, text=icon, font=self.fonts["icon"], bg=CARD_BG, fg=ICON).pack(side="left", padx=(0, px(8)))
        tk.Label(row, text=title, font=self.fonts["card_title"], bg=CARD_BG, fg=TEXT).pack(side="left")
        tk.Frame(parent, bg=DIVIDER, height=1).pack(fill="x", pady=(px(12), 0))
        return row

    def _build_ui(self) -> None:
        shadow = 12
        gutter = px(shadow)
        shell = tk.Frame(self.root, bg=WINDOW_BG)
        shell.pack(fill="both", expand=True, padx=px(14), pady=(px(16), px(8)))

        header = tk.Frame(shell, bg=WINDOW_BG)
        header.pack(fill="x", padx=gutter, pady=(0, px(4)))
        badge_size = px(42)
        badge = tk.Canvas(header, width=badge_size, height=badge_size, bg=WINDOW_BG, highlightthickness=0, bd=0)
        badge.pack(side="left")
        self._badge_photo = ImageTk.PhotoImage(
            render_round_rect(badge_size, badge_size, px(10), ACCENT, None, WINDOW_BG)
        )
        badge.create_image(0, 0, image=self._badge_photo, anchor="nw")
        badge.create_text(badge_size / 2, badge_size / 2 + 1, text=GLYPH_PICTURE, font=self.fonts["icon_badge"], fill="#FFFFFF")
        titles = tk.Frame(header, bg=WINDOW_BG)
        titles.pack(side="left", padx=(px(12), 0))
        tk.Label(titles, text="公众号素材整理", font=self.fonts["title"], bg=WINDOW_BG, fg=TEXT).pack(anchor="w")
        tk.Label(
            titles,
            text="粘贴文章链接，提取正文图片并导出 Word / PowerPoint",
            font=self.fonts["small"],
            bg=WINDOW_BG,
            fg=MUTED,
        ).pack(anchor="w")

        # ---- 导出设置卡片
        settings = RoundedBox(
            shell, parent_bg=WINDOW_BG, fill=CARD_BG, outline=BORDER, radius=10, padding=(20, 16), shadow=shadow
        )
        settings.pack(fill="x")
        body = settings.inner
        self._card_header(body, GLYPH_SETTINGS, "导出设置")

        self._section_label(body, "文章链接")
        self.url_field = InputField(
            body,
            parent_bg=CARD_BG,
            icon=GLYPH_LINK,
            placeholder="https://mp.weixin.qq.com/s/……",
            font=self.fonts["body"],
            icon_font=self.fonts["icon"],
        )
        self.url_field.pack(fill="x")
        self.url_field.add_action(
            lambda parent: IconButton(
                parent,
                text="粘贴",
                icon=GLYPH_PASTE,
                command=self._paste_url,
                parent_bg=CARD_BG,
                font=self.fonts["small"],
                icon_font=self.fonts["icon_small"],
                variant="ghost",
                height=28,
                padding_x=8,
            )
        )

        self._section_label(body, "文件名", "留空则使用文章标题")
        self.prefix_field = InputField(
            body,
            parent_bg=CARD_BG,
            icon=GLYPH_RENAME,
            placeholder="自动使用文章标题",
            font=self.fonts["body"],
            icon_font=self.fonts["icon"],
        )
        self.prefix_field.pack(fill="x")

        self._section_label(body, "保存位置", "每次导出新建一个子文件夹")
        location = tk.Frame(body, bg=CARD_BG)
        location.pack(fill="x")
        IconButton(
            location,
            text="更改",
            icon=GLYPH_FOLDER_OPEN,
            command=self._choose_path,
            parent_bg=CARD_BG,
            font=self.fonts["button"],
            icon_font=self.fonts["icon_small"],
            height=40,
            padding_x=14,
        ).pack(side="right", padx=(px(8), 0))
        self.path_box = RoundedBox(
            location, parent_bg=CARD_BG, fill=INFO_BG, outline=None, radius=8, padding=(12, 0), height=40
        )
        self.path_box.pack(side="left", fill="x", expand=True)
        tk.Label(self.path_box.inner, text=GLYPH_FOLDER, font=self.fonts["icon"], bg=INFO_BG, fg=ICON).pack(
            side="left", padx=(0, px(10))
        )
        self.path_label = tk.Label(self.path_box.inner, font=self.fonts["body"], bg=INFO_BG, anchor="w")
        self.path_label.pack(side="left", fill="x", expand=True)
        self.path_label.bind("<Configure>", lambda _event: self._refresh_path_label())
        self.path_box.bind("<Double-Button-1>", lambda _event: self._choose_path())
        bind_tree(self.path_box, "<Double-Button-1>", lambda _event: self._choose_path())

        self._section_label(body, "导出格式")
        tiles = tk.Frame(body, bg=CARD_BG)
        tiles.pack(fill="x")
        tiles.columnconfigure((0, 1), weight=1, uniform="tiles")
        self.word_tile = OptionTile(
            tiles,
            parent_bg=CARD_BG,
            title="Word 文档",
            subtitle="A4 竖版，图片按页宽排列",
            icon=GLYPH_DOCUMENT,
            variable=self.gen_word_var,
            fonts=self.fonts,
        )
        self.word_tile.grid(row=0, column=0, sticky="ew", padx=(0, px(6)))
        self.ppt_tile = OptionTile(
            tiles,
            parent_bg=CARD_BG,
            title="PowerPoint 演示文稿",
            subtitle="每页一张图，居中适配",
            icon=GLYPH_SLIDESHOW,
            variable=self.gen_ppt_var,
            fonts=self.fonts,
        )
        self.ppt_tile.grid(row=0, column=1, sticky="ew", padx=(px(6), 0))

        # ---- 操作按钮
        actions = tk.Frame(shell, bg=WINDOW_BG)
        actions.pack(fill="x", padx=gutter, pady=(0, px(2)))
        self.run_button = IconButton(
            actions,
            text="开始生成",
            icon=GLYPH_PLAY,
            command=self._start_task,
            parent_bg=WINDOW_BG,
            font=self.fonts["button_bold"],
            icon_font=self.fonts["icon_small"],
            variant="primary",
            height=42,
        )
        self.run_button.pack(side="left", fill="x", expand=True)
        self.cancel_button = IconButton(
            actions,
            text="取消",
            icon=GLYPH_CANCEL,
            command=self._cancel_task,
            parent_bg=WINDOW_BG,
            font=self.fonts["button"],
            icon_font=self.fonts["icon_small"],
            height=42,
        )
        self.cancel_button.pack(side="left", padx=(px(10), 0))
        self.cancel_button.set_enabled(False)
        self.open_button = IconButton(
            actions,
            text="打开输出文件夹",
            icon=GLYPH_FOLDER_OPEN,
            command=self._open_output_folder,
            parent_bg=WINDOW_BG,
            font=self.fonts["button"],
            icon_font=self.fonts["icon_small"],
            height=42,
        )
        self.open_button.pack(side="left", padx=(px(10), 0))
        self.open_button.set_enabled(False)

        # ---- 运行状态卡片
        status = RoundedBox(
            shell,
            parent_bg=WINDOW_BG,
            fill=CARD_BG,
            outline=BORDER,
            radius=10,
            padding=(20, 16),
            shadow=shadow,
            height=200,
        )
        status.pack(fill="both", expand=True)
        status_body = status.inner
        header_row = self._card_header(status_body, GLYPH_LIST, "运行状态")
        self.state_label = tk.Label(header_row, text="空闲", font=self.fonts["small"], bg=CARD_BG, fg=MUTED)
        self.state_label.pack(side="right")

        self.progress = ProgressBar(status_body, parent_bg=CARD_BG)
        self.progress.pack(fill="x", pady=(px(14), px(12)))

        log_box = RoundedBox(status_body, parent_bg=CARD_BG, fill=INFO_BG, radius=8, padding=(12, 8), height=80)
        log_box.pack(fill="both", expand=True)
        scrollbar = ttk.Scrollbar(log_box.inner, orient="vertical", style="Thin.Vertical.TScrollbar")

        def update_scrollbar(first: str, last: str) -> None:
            # 内容不足一屏时隐藏滚动条
            if float(first) <= 0 and float(last) >= 1:
                scrollbar.pack_forget()
            elif not scrollbar.winfo_ismapped():
                scrollbar.pack(side="right", fill="y", before=self.status_area)
            scrollbar.set(first, last)
        self.status_area = tk.Text(
            log_box.inner,
            bg=INFO_BG,
            fg=TEXT,
            relief="flat",
            bd=0,
            highlightthickness=0,
            font=self.fonts["log"],
            wrap="word",
            spacing1=px(2),
            spacing3=px(2),
            cursor="arrow",
            selectbackground=SELECTION_BG,
            selectforeground=TEXT,
            inactiveselectbackground=SELECTION_BG,
            yscrollcommand=update_scrollbar,
            state="disabled",
        )
        self.status_area.pack(side="left", fill="both", expand=True)
        scrollbar.configure(command=self.status_area.yview)
        self.status_area.tag_configure("time", foreground=PLACEHOLDER)
        self.status_area.tag_configure("placeholder", foreground=PLACEHOLDER)
        for level in LOG_COLORS:
            self.status_area.tag_configure(
                f"icon_{level}", foreground=LOG_ICON_COLORS[level], font=self.fonts["icon_small"]
            )
            self.status_area.tag_configure(f"msg_{level}", foreground=LOG_COLORS[level])
        self._show_log_placeholder()

    def _bind_shortcuts(self) -> None:
        for field in (self.url_field, self.prefix_field):
            field.entry.bind("<Return>", lambda _event: self._start_task())
        self.root.bind("<Escape>", lambda _event: self._cancel_task())

    # ---------------------------------------------------------------- 交互

    def _refresh_path_label(self) -> None:
        if self.save_dir:
            width = self.path_label.winfo_width() - px(4)
            text = elide_middle(self.save_dir, self.fonts["body"], width)
            self.path_label.configure(text=text, fg=TEXT)
        else:
            self.path_label.configure(text="尚未选择文件夹，点击右侧“更改”", fg=PLACEHOLDER)

    def _paste_url(self) -> None:
        try:
            text = self.root.clipboard_get().strip()
        except tk.TclError:
            return
        # 从分享文案中提取第一个链接
        match = re.search(r"https?://\S+", text)
        self.url_field.var.set(match.group(0) if match else text)
        self.url_field.entry.icursor("end")
        self.url_field.entry.focus_set()

    def _choose_path(self) -> None:
        if self.is_running:
            return
        folder = filedialog.askdirectory(initialdir=self.save_dir or None, title="选择保存位置")
        if folder:
            self.save_dir = os.path.normpath(folder)
            self.path_box.set_style(outline=None)
            self._refresh_path_label()
            self._save_settings()

    def _open_output_folder(self) -> None:
        if self.last_output_folder and os.path.isdir(self.last_output_folder):
            open_in_file_manager(self.last_output_folder)

    def _cancel_task(self) -> None:
        if not self.is_running or self.cancel_event.is_set():
            return
        self.cancel_event.set()
        self.cancel_button.set_enabled(False)
        self._set_state_text("正在取消…", WARNING)

    def _save_settings(self) -> None:
        save_config(
            {
                "save_dir": self.save_dir,
                "want_word": bool(self.gen_word_var.get()),
                "want_ppt": bool(self.gen_ppt_var.get()),
            }
        )

    def _on_close(self) -> None:
        if self.is_running:
            self.cancel_event.set()
        self._save_settings()
        self.root.destroy()

    def _set_state_text(self, text: str, color: str = MUTED) -> None:
        self.state_label.configure(text=text, fg=color)

    def _show_log_placeholder(self) -> None:
        self.status_area.configure(state="normal")
        self.status_area.delete("1.0", "end")
        self.status_area.insert("end", "粘贴文章链接后点击「开始生成」，处理过程会显示在这里。", "placeholder")
        self.status_area.configure(state="disabled")
        self.log_is_placeholder = True

    def _clear_log(self) -> None:
        self.status_area.configure(state="normal")
        self.status_area.delete("1.0", "end")
        self.status_area.configure(state="disabled")
        self.log_is_placeholder = False

    def _append_status(self, message: str, level: str) -> None:
        if level not in LOG_COLORS:
            level = "info"
        if self.log_is_placeholder:
            self._clear_log()
        timestamp = datetime.datetime.now().strftime("%H:%M:%S")
        self.status_area.configure(state="normal")
        self.status_area.insert("end", LOG_GLYPHS[level], f"icon_{level}")
        self.status_area.insert("end", f"  {timestamp}  ", "time")
        self.status_area.insert("end", f"{message}\n", f"msg_{level}")
        self.status_area.see("end")
        self.status_area.configure(state="disabled")

    def _set_running_state(self, running: bool) -> None:
        self.is_running = running
        self.run_button.set_enabled(not running)
        self.run_button.set_text("正在处理…" if running else "开始生成")
        self.cancel_button.set_enabled(running)
        for field in (self.url_field, self.prefix_field):
            field.entry.configure(state="disabled" if running else "normal", disabledbackground=CARD_BG,
                                  disabledforeground=MUTED)

    def _finish(self, outcome: str, folder: str, elapsed: float) -> None:
        self._set_running_state(False)
        if folder and os.path.isdir(folder):
            self.last_output_folder = folder
            self.open_button.set_enabled(True)
        if outcome == "success":
            self.progress.set(1.0, SUCCESS)
            self._set_state_text(f"已完成 · 用时 {elapsed:.1f} 秒", SUCCESS)
        elif outcome == "cancelled":
            self._set_state_text("已取消", WARNING)
        else:
            self.progress.set(self.progress.value, DANGER)
            self._set_state_text("未完成", DANGER)

    def _poll_status_queue(self) -> None:
        try:
            while True:
                item = self.status_queue.get_nowait()
                kind = item.get("type")
                if kind == "progress":
                    self.progress.set(float(item.get("value", 0)))
                elif kind == "stage":
                    self._set_state_text(str(item.get("label", "")))
                elif kind == "log":
                    self._append_status(str(item["message"]), str(item["level"]))
                elif kind == "done":
                    self._finish(str(item["outcome"]), str(item["folder"]), float(item["elapsed"]))
        except queue.Empty:
            pass
        try:
            self.root.after(100, self._poll_status_queue)
        except tk.TclError:
            pass

    def _validate_inputs(self) -> tuple[str, str] | None:
        article_url = self.url_field.get()
        prefix = self.prefix_field.get()
        problems: list[str] = []

        if not is_http_url(article_url):
            self.url_field.set_error(True)
            problems.append("请填写有效的文章链接")
        if not self.save_dir or not os.path.isdir(self.save_dir):
            self.path_box.set_style(outline=DANGER)
            problems.append("请选择保存位置")
        if not self.gen_word_var.get() and not self.gen_ppt_var.get():
            self.word_tile.set_error(True)
            self.ppt_tile.set_error(True)
            problems.append("请至少选择一种导出格式")
        if problems:
            self._set_state_text("，".join(problems), DANGER)
            if not is_http_url(article_url):
                self.url_field.entry.focus_set()
            return None
        return article_url, prefix

    def _start_task(self) -> None:
        if self.is_running:
            return

        validated = self._validate_inputs()
        if not validated:
            return

        article_url, prefix = validated
        self._save_settings()
        self.cancel_event = threading.Event()
        self.status_queue = queue.Queue()
        self.open_button.set_enabled(False)
        self.progress.set(0)
        self._clear_log()
        self._set_state_text("准备中")
        self._set_running_state(True)
        threading.Thread(
            target=run_export,
            args=(
                article_url,
                prefix,
                self.save_dir,
                bool(self.gen_word_var.get()),
                bool(self.gen_ppt_var.get()),
                self.status_queue,
                self.cancel_event,
            ),
            daemon=True,
        ).start()


def enable_high_dpi() -> None:
    if not sys.platform.startswith("win"):
        return
    try:
        import ctypes

        ctypes.windll.shcore.SetProcessDpiAwareness(1)
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("ys199.WeixinMaterialTool")
    except (AttributeError, OSError):
        pass


def main() -> None:
    global UI_SCALE
    enable_high_dpi()
    root = tk.Tk()
    UI_SCALE = max(1.0, root.winfo_fpixels("1i") / 96)
    MaterialApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
